# coding=gb2312
import os
from itertools import groupby, chain

from openpyxl import Workbook
from openpyxl import load_workbook

from common.config import config
from common.data_types import TestCase
from common.files import *
from utils.io import csv2, excel
from utils.io.files import add_bom_utf8
from utils.parsers import mm
from utils.parsers import xm


def get_title():
    return [field for field in config.get_item("fields")]


def _iter_merged_by_step(raw_cases):
    """按案例步骤合并，一个步骤对应多个结果时自动合并，同时更新stepn"""
    for title, cases in groupby(raw_cases, key=(lambda case: case.title)):
        for n, (step, cases_) in enumerate(groupby(cases, key=(lambda case: case.step)), 1):
            result = "\n".join((c.result for c in cases_))
            yield TestCase(title, step, result, n)


def _iter_merged_by_title(merged_cases, sep='_'):
    """按案例标题合并"""

    common_fields = config.get_item("common_fields")
    import_path = common_fields.get("导入路径", "")
    author = common_fields.get("设计者", "")
    component = common_fields.get("归属系统", "")
    fields = config.get_item("fields")
    precondition = fields.get("Precondition", "")
    status = fields.get("Status", "Approved")
    priority = fields.get("Priority", "High")
    estimated_time = fields.get("Estimated Time", "")
    labels = fields.get("Labels", "正例")
    coverage = fields.get("Coverage(Issues)", "")
    merge_steps = config.get_item("output_format")["current_item"]
    yield get_title()
    if merge_steps == "合并步骤为一行":
        for title, cases in groupby(merged_cases, key=(lambda case: case.title)):
            cases = tuple(cases)
            step = "\n".join((case.step for case in cases))
            result = "\n".join((case.result for case in cases))
            title = sep.join(title)
            yield (
                title, title,
                precondition, import_path, status, priority, component,
                author, estimated_time, labels, coverage,
                step, result)

    else:
        if merge_steps == "不合并步骤":
            for c in merged_cases:
                if c.stepn == 1:
                    title = sep.join(c.title)
                    step = c.step
                    result = c.result
                    yield (title, title,
                           precondition, import_path, status, priority, component,
                           author, estimated_time, labels, coverage,
                           step, result)
                else:
                    yield (
                        "", "",
                        "", "", "", "", "",
                        "", "", "", "",
                        c.step, c.result)

        else:
            for case in merged_cases:
                title = sep.join(chain(case.title, (case.step, case.result)))
                step = sep.join(chain(case.title[1:], (case.step,)))
                result = case.result
                yield (
                    title, title,
                    precondition, import_path, status, priority, component,
                    author, estimated_time, labels, coverage,
                    step, result)


def case_write_file():
    pass


def get_template_header(worksheet):
    headers = []
    for cell in next(worksheet.iter_rows(min_row=1, max_row=1)):
        headers.append(cell.value)

    yield headers


def set_sheet_header(worksheet, row):
    worksheet.append(next(row))


def set_sheet_row(worksheet, row):
    worksheet.append(row)


def merge_cases(input_file, cases):
    common_fields = config.get_item("common_fields")
    input_file_path = Path(input_file)
    output_file_path = input_file_path.parent.joinpath(input_file_path.stem).with_suffix('.xlsx')
    temp_path = get_dirs_realpath().parent.joinpath("testcase_template")
    temp_file_path = get_dirs_realpath(temp_path)
    temp_workbook = load_workbook(filename=temp_file_path)
    temp_worksheet = temp_workbook.active
    headers = get_template_header(temp_worksheet)
    output_workbook = Workbook()
    output_worksheet = output_workbook.active
    output_worksheet.title = "测试用例"
    set_sheet_header(output_worksheet, headers)
    for case in cases:
        case.us_name = ""
        case.no = ""
        case.path = common_fields.get("导入路径", "")
        case.titles = "_".join(chain(case.title, (case.step, case.result)))
        case.desc = ""
        case.pre_step = ""
        case.steps = "_".join(chain(case.title[1:], (case.step,)))
        case.results = case.result
        case.real_result = ""
        case.level = common_fields.get("优先级", "")
        case.auther = common_fields.get("设计者", "")
        case.type = common_fields.get("类型", "")
        case.app = common_fields.get("归属系统", "")
        case.label = ""
        del case.stepn
        del case.step
        del case.title
        del case.result
        case_row = case.__dict__.values()
        set_sheet_row(output_worksheet, tuple(case_row))
    output_workbook.save(output_file_path)
    return output_file_path


def to_csv(file, data):
    csv2.write(file, data)
    add_bom_utf8(file)
    return file


def to_excel(file, data):
    excel.write(file, data, sheet_name="测试用例")
    return file


def build(file, output_format=0, file_type='excel'):
    if file.lower().endswith(".mm"):
        parser = mm.parse_freemind
    else:
        if file.lower().endswith(".xmind"):
            parser = xm.parse_xm
        else:
            raise ValueError("入参应该是str：.mm或.xmind文件")
    raw_cases = parser(file)
    if output_format == "仅输出标题":
        return str(merge_cases(file, raw_cases))
    else:
        pre_cases = _iter_merged_by_step(raw_cases)
        final_cases = _iter_merged_by_title(pre_cases)
    output_file, _ = os.path.splitext(file)
    if file_type == "excel":
        to_file = to_excel
        output_file += ".xls"
    else:
        to_file = to_csv
        output_file += ".csv"
    return to_file(output_file, final_cases)
